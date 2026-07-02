import io
import zipfile
import openpyxl
from collections import Counter
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Font, PatternFill
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from ..models import Term, SchoolClass, Student, Assessment, Grade, Attendance, ReportCard, Subject, AcademicYear
from ..decorators import admin_required
from ..utils import get_term_subject_mark, get_grade_letter

@login_required
@admin_required
def reports_menu(request):
    return render(request, 'reports/reports_menu.html')


@login_required
@admin_required
def generate_report_cards(request):
    context = {
        'terms':    Term.objects.all().order_by('-start_date'),
        'classes':  SchoolClass.objects.all().order_by('name'),
        'students': Student.objects.filter(is_active=True).select_related('user').prefetch_related('enrollments', 'enrollments__class_assigned'),
    }
    return render(request, 'reports/generate_reports.html', context)


def generate_single_report_card_pdf(request, student, term):
    enrollment = student.enrollments.filter(academic_year=term.academic_year).first()
    target_class = enrollment.class_assigned if enrollment else None

    grades = []
    if target_class:
        subject_ids = Assessment.objects.filter(
            class_assigned=target_class,
            term=term,
        ).values_list('subject_id', flat=True).distinct()
        for subject_id in subject_ids:
            subject = Subject.objects.get(pk=subject_id)
            mark = get_term_subject_mark(student, subject, term)
            if mark is not None:
                grades.append({
                    'subject': subject,
                    'average_score': mark,
                    'grade_letter': get_grade_letter(mark),
                    'teacher_comment': '',
                })
        grades.sort(key=lambda g: g['subject'].name)

    if not grades:
        messages.warning(request, 'No grades found for this student in the selected term.')
        return redirect('generate_report_cards')

    overall_average = round(sum(g['average_score'] for g in grades) / len(grades), 1)
    overall_grade   = get_grade_letter(overall_average)

    att = Attendance.objects.filter(
        student=student, date__gte=term.start_date, date__lte=term.end_date
    )
    total_days          = att.count()
    days_present        = att.filter(status='present').count()
    days_absent         = att.filter(status='absent').count()
    days_late           = att.filter(status='late').count()
    attendance_pct      = round((days_present / total_days * 100), 1) if total_days else 0

    buffer = io.BytesIO()
    doc    = SimpleDocTemplate(buffer, pagesize=A4,
                               rightMargin=0.5*inch, leftMargin=0.5*inch,
                               topMargin=0.5*inch, bottomMargin=0.5*inch)
    elements = []
    styles   = getSampleStyleSheet()

    title_style = ParagraphStyle('Title', parent=styles['Heading1'],
                                 fontSize=18, alignment=TA_CENTER, spaceAfter=6)
    h2_style    = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=12, spaceAfter=6)

    elements.append(Paragraph("YOUR SCHOOL NAME", title_style))
    elements.append(Paragraph("School Address | Phone | Email", styles['Normal']))
    elements.append(Spacer(1, 0.2*inch))
    elements.append(Paragraph("STUDENT REPORT CARD", h2_style))
    elements.append(Paragraph(f"{term.name} — {term.academic_year.name}", styles['Normal']))
    elements.append(Spacer(1, 0.3*inch))

    student_class_name = target_class.name if target_class else 'N/A'

    student_table = Table([
        ['Student Name:', student.user.get_full_name(), 'Student ID:', student.student_id_number],
        ['Class:', student_class_name, 'Academic Year:', term.academic_year.name],
        ['Term:', term.name, 'Report Date:', timezone.now().date().strftime('%Y-%m-%d')],
    ], colWidths=[1.5*inch, 2.5*inch, 1.5*inch, 2*inch])
    student_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
        ('BACKGROUND', (2, 0), (2, -1), colors.lightgrey),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(student_table)
    elements.append(Spacer(1, 0.3*inch))
    elements.append(Paragraph("Academic Performance", h2_style))

    grade_data = [['Subject', 'Score (%)', 'Grade', 'Teacher Comment']]
    for g in grades:
        grade_data.append([g['subject'].name, f"{g['average_score']}%", g['grade_letter'], g['teacher_comment'] or '-'])

    grade_table = Table(grade_data, colWidths=[2*inch, 1*inch, 1*inch, 3.5*inch])
    grade_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#333333')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9f9f9')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(grade_table)
    elements.append(Spacer(1, 0.2*inch))

    summary_table = Table([
        ['Overall Average:', f"{overall_average}%", 'Overall Grade:', overall_grade],
        ['Total Subjects:', str(len(grades)), '', ''],
    ], colWidths=[1.5*inch]*4)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f0f0f0')),
        ('BOX', (0, 0), (-1, -1), 2, colors.HexColor('#333333')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.3*inch))
    elements.append(Paragraph("Attendance Record", h2_style))

    att_table = Table([
        ['Days Present:', str(days_present), 'Days Absent:', str(days_absent)],
        ['Days Late:', str(days_late), 'Attendance Rate:', f"{attendance_pct}%"],
    ], colWidths=[1.5*inch]*4)
    att_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
        ('BACKGROUND', (2, 0), (2, -1), colors.lightgrey),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(att_table)
    elements.append(Spacer(1, 0.3*inch))

    report_card = ReportCard.objects.filter(student=student, term=term).first()
    if report_card and report_card.class_teacher_comment:
        elements.append(Paragraph("Class Teacher's Comment:", h2_style))
        elements.append(Paragraph(report_card.class_teacher_comment, styles['Normal']))
        elements.append(Spacer(1, 0.2*inch))
    if report_card and report_card.principal_comment:
        elements.append(Paragraph("Principal's Comment:", h2_style))
        elements.append(Paragraph(report_card.principal_comment, styles['Normal']))
        elements.append(Spacer(1, 0.2*inch))

    elements.append(Spacer(1, 0.4*inch))
    ct_name = (
        target_class.class_teacher.user.get_full_name()
        if target_class and target_class.class_teacher
        else "Not Assigned"
    )
    sig_table = Table([
        ['_________________________', '_________________________'],
        [f'Class Teacher: {ct_name}', 'Principal'],
    ], colWidths=[3.5*inch, 3.5*inch])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('TOPPADDING', (0, 1), (-1, 1), 10),
    ]))
    elements.append(sig_table)
    elements.append(Spacer(1, 0.3*inch))
    elements.append(Paragraph(
        f"Official document — YOUR SCHOOL NAME<br/>Generated on {timezone.now().date()}",
        styles['Normal'],
    ))

    doc.build(elements)
    pdf_data = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf_data, content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="report_card_{student.student_id_number}_{term.name}.pdf"'
    )
    return response


@login_required
def preview_report_card(request):
    term_id    = request.GET.get('term')
    class_id   = request.GET.get('class')
    student_id = request.GET.get('student')

    if not term_id or not class_id:
        messages.error(request, 'Please select term and class.')
        return redirect('generate_report_cards')

    term = get_object_or_404(Term, pk=term_id)

    if student_id:
        student = get_object_or_404(Student, pk=student_id)
    else:
        class_obj = get_object_or_404(SchoolClass, pk=class_id)
        student   = Student.objects.filter(
            enrollments__class_assigned=class_obj,
            enrollments__is_active=True,
            is_active=True
        ).first()
        if not student:
            messages.error(request, 'No students found in this class.')
            return redirect('generate_report_cards')

    return generate_single_report_card_pdf(request, student, term)


@login_required
@admin_required
def bulk_generate_reports(request):
    if request.method != 'POST':
        return redirect('generate_report_cards')

    term_id  = request.POST.get('term')
    class_id = request.POST.get('class')
    if not term_id or not class_id:
        messages.error(request, 'Please select term and class.')
        return redirect('generate_report_cards')

    term      = get_object_or_404(Term, pk=term_id)
    class_obj = get_object_or_404(SchoolClass, pk=class_id)
    students  = Student.objects.filter(
        enrollments__class_assigned=class_obj,
        enrollments__is_active=True,
        is_active=True
    ).distinct()

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        for student in students:
            has_grades = Grade.objects.filter(
                student=student, assessment__term=term
            ).exists()
            if not has_grades:
                continue
            pdf_response = generate_single_report_card_pdf(request, student, term)
            filename = f"{student.student_id_number}_{student.user.last_name}_{student.user.first_name}.pdf"
            zf.writestr(filename, pdf_response.content)

    zip_buffer.seek(0)
    response = HttpResponse(zip_buffer.read(), content_type='application/zip')
    response['Content-Disposition'] = (
        f'attachment; filename="report_cards_{class_obj.name}_{term.name}.zip"'
    )
    return response


@login_required
@admin_required
def grade_reports(request):
    terms    = Term.objects.all().order_by('-start_date')
    selected_term    = request.GET.get('term', terms.first().id if terms.exists() else None)
    selected_class   = request.GET.get('class', '')
    selected_subject = request.GET.get('subject', '')

    term = get_object_or_404(Term, pk=selected_term) if selected_term else None
    students = Student.objects.filter(is_active=True)
    if selected_class:
        students = students.filter(
            enrollments__class_assigned_id=selected_class,
            enrollments__is_active=True
        ).distinct()

    grade_summaries = []
    if term:
        if selected_subject:
            subject = get_object_or_404(Subject, pk=selected_subject)
            for student in students:
                mark = get_term_subject_mark(student, subject, term)
                if mark is not None:
                    grade_summaries.append({
                        'student': student,
                        'subject': subject,
                        'average_score': mark,
                        'grade_letter': get_grade_letter(mark),
                    })
        else:
            for student in students:
                if student.current_class:
                    subject_ids = Assessment.objects.filter(
                        term=term,
                        class_assigned=student.current_class,
                    ).values_list('subject_id', flat=True).distinct()
                    marks = []
                    for subject_id in subject_ids:
                        subject = Subject.objects.get(pk=subject_id)
                        mark = get_term_subject_mark(student, subject, term)
                        if mark is not None:
                            marks.append(mark)
                    if marks:
                        avg = round(sum(marks) / len(marks), 1)
                        grade_summaries.append({
                            'student': student,
                            'average_score': avg,
                            'grade_letter': get_grade_letter(avg),
                        })

    if grade_summaries:
        scores = [g['average_score'] for g in grade_summaries]
        stats = {
            'total_students': len(set(g['student'].id for g in grade_summaries)),
            'average_score':  round(sum(scores) / len(scores), 1),
            'highest_score':  round(max(scores), 1),
            'lowest_score':   round(min(scores), 1),
        }
    else:
        stats = {'total_students': 0, 'average_score': 0, 'highest_score': 0, 'lowest_score': 0}

    grade_letters    = [g['grade_letter'] for g in grade_summaries]
    grade_counts     = Counter(grade_letters)
    total            = len(grade_letters) or 1
    grade_distribution = [
        {'grade': g, 'count': grade_counts.get(g, 0),
         'percentage': round(grade_counts.get(g, 0) / total * 100, 1)}
        for g in ['A+', 'A', 'B+', 'B', 'C+', 'C', 'D+', 'D', 'F']
    ]

    top_students = []
    selected_subject_name = ''
    if selected_subject:
        top_students = sorted(grade_summaries, key=lambda x: x['average_score'], reverse=True)[:10]
        selected_subject_name = get_object_or_404(Subject, pk=selected_subject).name

    context = {
        'terms': terms,
        'selected_term': str(selected_term),
        'selected_class': selected_class,
        'selected_subject': selected_subject,
        'selected_subject_name': selected_subject_name,
        'classes': SchoolClass.objects.all(),
        'subjects': Subject.objects.all(),
        'stats': stats,
        'grade_distribution': grade_distribution,
        'top_students': top_students,
    }
    return render(request, 'reports/grade_reports.html', context)


@login_required
@admin_required
def export_grade_report(request):
    term_id    = request.GET.get('term')
    class_id   = request.GET.get('class', '')
    subject_id = request.GET.get('subject', '')
    term       = get_object_or_404(Term, pk=term_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Grade Report"
    ws['A1'] = "Grade Report"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Term: {term.name}"
    ws.append([])
    ws.append(['Student ID', 'Name', 'Class', 'Subject', 'Score (%)', 'Grade'])
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    students = Student.objects.filter(is_active=True)
    if class_id:
        students = students.filter(
            enrollments__class_assigned_id=class_id,
            enrollments__is_active=True
        ).distinct()

    if subject_id:
        subject = get_object_or_404(Subject, pk=subject_id)
        for student in students:
            target_class = student.get_class_for_term(term)
            mark = get_term_subject_mark(student, subject, term)
            if mark is not None:
                ws.append([
                    student.student_id_number,
                    student.user.get_full_name(),
                    target_class.name if target_class else 'N/A',
                    subject.name,
                    mark,
                    get_grade_letter(mark),
                ])
    else:
        for student in students:
            target_class = student.get_class_for_term(term)
            if target_class:
                subject_ids = Assessment.objects.filter(
                    term=term,
                    class_assigned=target_class,
                ).values_list('subject_id', flat=True).distinct()
                for subject_id in subject_ids:
                    subject = Subject.objects.get(pk=subject_id)
                    mark = get_term_subject_mark(student, subject, term)
                    if mark is not None:
                        ws.append([
                            student.student_id_number,
                            student.user.get_full_name(),
                            target_class.name if target_class else 'N/A',
                            subject.name,
                            mark,
                            get_grade_letter(mark),
                        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=grade_report_{term.name}.xlsx'
    wb.save(response)
    return response


@login_required
@admin_required
def export_student_list(request):
    class_id = request.GET.get('class', '')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student List"
    ws.append(['Student ID', 'First Name', 'Last Name', 'Email', 'Gender',
               'Class', 'Emergency Contact', 'Phone', 'Status'])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    students = Student.objects.select_related('user').prefetch_related('enrollments', 'enrollments__class_assigned')
    if class_id:
        students = students.filter(
            enrollments__class_assigned_id=class_id,
            enrollments__is_active=True
        ).distinct()

    for s in students:
        ws.append([
            s.student_id_number,
            s.user.first_name,
            s.user.last_name,
            s.user.email,
            s.get_gender_display(),
            s.current_class.name if s.current_class else 'N/A',
            s.emergency_contact_name,
            s.emergency_contact_phone,
            'Active' if s.is_active else 'Inactive',
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=student_list.xlsx'
    wb.save(response)
    return response
