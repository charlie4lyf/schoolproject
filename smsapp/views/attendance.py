import openpyxl
from datetime import datetime, timedelta
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill
from ..models import SchoolClass, Attendance, Student, Term, ClassSubjectTeacher, ParentStudent
from ..decorators import admin_required

@login_required
def mark_attendance(request, class_id):
    class_obj = get_object_or_404(SchoolClass, pk=class_id)

    if request.user.role == 'teacher':
        teacher = request.user.teacher_profile
        if not ClassSubjectTeacher.objects.filter(teacher=teacher, class_assigned=class_obj).exists():
            messages.error(request, 'You do not have permission to mark attendance for this class.')
            return redirect('teacher_dashboard')
    elif request.user.role != 'admin':
        messages.error(request, 'You do not have permission to access this page.')
        return redirect('dashboard')

    selected_date = request.GET.get('date', timezone.now().date().isoformat())
    selected_date = datetime.strptime(selected_date, '%Y-%m-%d').date()
    today = timezone.now().date()
    can_edit = selected_date >= today or request.user.role == 'admin'

    students = Student.objects.filter(
        enrollments__class_assigned=class_obj,
        enrollments__is_active=True,
        is_active=True,
    ).select_related('user').distinct().order_by('user__last_name', 'user__first_name')

    attendance_records = Attendance.objects.filter(class_assigned=class_obj, date=selected_date)
    existing_attendance = {
        r.student_id: {'status': r.status, 'remarks': r.remarks}
        for r in attendance_records
    }
    already_marked = attendance_records.exists()

    for student in students:
        info = existing_attendance.get(student.id, {})
        student.attendance_status  = info.get('status', 'present')
        student.attendance_remarks = info.get('remarks', '')

    if request.method == 'POST' and can_edit:
        marked_by = request.user.teacher_profile if request.user.role == 'teacher' else None
        current_term = Term.objects.filter(is_active=True).first()

        for student in students:
            status  = request.POST.get(f'status_{student.id}', 'present')
            remarks = request.POST.get(f'remarks_{student.id}', '')

            defaults = {
                'class_assigned': class_obj,
                'status':         status,
                'remarks':        remarks,
                'marked_by':      marked_by,
            }
            if current_term:
                defaults['term'] = current_term

            Attendance.objects.update_or_create(
                student=student,
                date=selected_date,
                defaults=defaults,
            )

        messages.success(request, f'Attendance marked for {class_obj.name} on {selected_date}')
        return redirect('teacher_dashboard' if request.user.role == 'teacher' else 'admin_dashboard')

    context = {
        'class': class_obj,
        'students': students,
        'selected_date': selected_date,
        'today': today,
        'can_edit': can_edit,
        'already_marked': already_marked,
    }
    return render(request, 'attendance/mark_attendance.html', context)


@login_required
def attendance_history(request):
    if request.user.role not in ['admin', 'teacher']:
        messages.error(request, 'You do not have permission to view this attendance page.')
        if request.user.role == 'student':
            return redirect('student_attendance_report', student_id=request.user.student_profile.id)
        return redirect('parent_dashboard')

    selected_class = request.GET.get('class', '')
    start_date = datetime.strptime(
        request.GET.get('start_date', (timezone.now().date() - timedelta(days=30)).isoformat()),
        '%Y-%m-%d'
    ).date()
    end_date = datetime.strptime(
        request.GET.get('end_date', timezone.now().date().isoformat()),
        '%Y-%m-%d'
    ).date()

    if request.user.role == 'teacher':
        teacher = request.user.teacher_profile
        classes = SchoolClass.objects.filter(subject_assignments__teacher=teacher).distinct()
    else:
        classes = SchoolClass.objects.all()

    attendance_query = Attendance.objects.filter(date__gte=start_date, date__lte=end_date)
    if selected_class:
        attendance_query = attendance_query.filter(class_assigned_id=selected_class)
    elif request.user.role == 'teacher':
        attendance_query = attendance_query.filter(
            class_assigned_id__in=classes.values_list('id', flat=True)
        )

    attendance_by_date = {}
    current = start_date
    while current <= end_date:
        day_records = attendance_query.filter(date=current)
        if day_records.exists():
            total   = day_records.count()
            present = day_records.filter(status='present').count()
            attendance_by_date[current] = {
                'is_marked':     True,
                'present_count': present,
                'absent_count':  day_records.filter(status='absent').count(),
                'late_count':    day_records.filter(status='late').count(),
                'excused_count': day_records.filter(status='excused').count(),
                'percentage':    round((present / total * 100), 1) if total else 0,
            }
        else:
            attendance_by_date[current] = {'is_marked': False}
        current += timedelta(days=1)

    context = {
        'classes': classes,
        'selected_class': selected_class,
        'start_date': start_date,
        'end_date': end_date,
        'attendance_by_date': attendance_by_date,
    }
    return render(request, 'attendance/attendance_history.html', context)


@login_required
def attendance_by_date(request):
    if request.user.role not in ['admin', 'teacher']:
        messages.error(request, 'You do not have permission to view this attendance page.')
        if request.user.role == 'student':
            return redirect('student_attendance_report', student_id=request.user.student_profile.id)
        return redirect('parent_dashboard')

    selected_date = datetime.strptime(
        request.GET.get('date', timezone.now().date().isoformat()), '%Y-%m-%d'
    ).date()
    selected_class = request.GET.get('class', '')

    attendance_records = Attendance.objects.filter(
        date=selected_date
    ).select_related('student__user', 'class_assigned')

    if selected_class:
        attendance_records = attendance_records.filter(class_assigned_id=selected_class)

    total   = attendance_records.count()
    present = attendance_records.filter(status='present').count()

    context = {
        'selected_date': selected_date,
        'selected_class': selected_class,
        'classes': SchoolClass.objects.all(),
        'attendance_records': attendance_records,
        'summary': {
            'total':      total,
            'present':    present,
            'absent':     attendance_records.filter(status='absent').count(),
            'late':       attendance_records.filter(status='late').count(),
            'excused':    attendance_records.filter(status='excused').count(),
            'percentage': round((present / total * 100), 1) if total else 0,
        },
    }
    return render(request, 'attendance/attendance_by_date.html', context)


@login_required
def student_attendance_report(request, student_id):
    student = get_object_or_404(Student.objects.select_related('user'), pk=student_id)

    if request.user.role == 'student':
        if request.user.student_profile.id != student.id:
            messages.error(request, 'You can only view your own attendance.')
            return redirect('student_dashboard')
    elif request.user.role == 'parent':
        if not ParentStudent.objects.filter(
            parent=request.user.parent_profile, student=student
        ).exists():
            messages.error(request, "You can only view your children's attendance.")
            return redirect('parent_dashboard')

    end_date   = timezone.now().date()
    start_date = datetime.strptime(
        request.GET.get('start_date', (end_date - timedelta(days=90)).isoformat()), '%Y-%m-%d'
    ).date()
    end_date = datetime.strptime(
        request.GET.get('end_date', end_date.isoformat()), '%Y-%m-%d'
    ).date()

    attendance_records = Attendance.objects.filter(
        student=student, date__gte=start_date, date__lte=end_date
    ).order_by('-date')

    total_days    = attendance_records.count()
    present_count = attendance_records.filter(status__in=['present', 'late']).count()
    absent_count  = attendance_records.filter(status='absent').count()
    percentage    = round((present_count / total_days * 100), 1) if total_days else 0

    consecutive_absences = 0
    for record in attendance_records[:10]:
        if record.status == 'absent':
            consecutive_absences += 1
        else:
            break

    context = {
        'student': student,
        'attendance_records': attendance_records,
        'stats': {
            'total_days':           total_days,
            'present_count':        present_count,
            'absent_count':         absent_count,
            'percentage':           percentage,
            'consecutive_absences': consecutive_absences,
        },
        'start_date': start_date,
        'end_date':   end_date,
    }
    return render(request, 'attendance/student_attendance.html', context)


@login_required
@admin_required
def class_attendance_report(request, class_id):
    class_obj = get_object_or_404(SchoolClass, pk=class_id)
    end_date   = timezone.now().date()
    start_date = datetime.strptime(
        request.GET.get('start_date', (end_date - timedelta(days=30)).isoformat()), '%Y-%m-%d'
    ).date()
    end_date = datetime.strptime(
        request.GET.get('end_date', end_date.isoformat()), '%Y-%m-%d'
    ).date()

    students = Student.objects.filter(
        enrollments__class_assigned=class_obj,
        enrollments__is_active=True,
        is_active=True
    ).distinct()
    student_stats, alerts = [], []

    for student in students:
        records    = Attendance.objects.filter(student=student, date__gte=start_date, date__lte=end_date)
        total_days = records.count()
        present    = records.filter(status__in=['present', 'late']).count()
        absent     = records.filter(status='absent').count()
        late       = records.filter(status='late').count()
        percentage = round((present / total_days * 100), 1) if total_days else 0

        consecutive = 0
        for r in Attendance.objects.filter(student=student).order_by('-date')[:10]:
            if r.status == 'absent':
                consecutive += 1
            else:
                break

        student_stats.append({
            'student': student, 'total_days': total_days,
            'present_count': present, 'absent_count': absent,
            'late_count': late, 'percentage': percentage,
            'consecutive_absences': consecutive,
        })

        if consecutive >= 3:
            alerts.append(f"{student.user.get_full_name()} absent for {consecutive} consecutive days")
        if percentage < 85 and total_days > 0:
            alerts.append(f"{student.user.get_full_name()} has low attendance ({percentage}%)")

    context = {
        'class': class_obj,
        'student_stats': student_stats,
        'start_date': start_date,
        'end_date': end_date,
        'alerts': alerts,
    }
    return render(request, 'attendance/class_attendance_report.html', context)


@login_required
@admin_required
def export_class_attendance(request, class_id):
    class_obj  = get_object_or_404(SchoolClass, pk=class_id)
    end_date   = timezone.now().date()
    start_date = datetime.strptime(
        request.GET.get('start_date', (end_date - timedelta(days=30)).isoformat()), '%Y-%m-%d'
    ).date()
    end_date = datetime.strptime(
        request.GET.get('end_date', end_date.isoformat()), '%Y-%m-%d'
    ).date()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    ws['A1'] = f"Attendance Report — {class_obj.name}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Period: {start_date} to {end_date}"
    ws.append([])
    ws.append(['Student ID', 'Student Name', 'Total Days', 'Present', 'Absent', 'Late', 'Percentage'])
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    for student in Student.objects.filter(
        enrollments__class_assigned=class_obj,
        enrollments__is_active=True,
        is_active=True
    ).distinct():
        records    = Attendance.objects.filter(student=student, date__gte=start_date, date__lte=end_date)
        total      = records.count()
        present    = records.filter(status__in=['present', 'late']).count()
        percentage = round((present / total * 100), 1) if total else 0
        ws.append([
            student.student_id_number,
            student.user.get_full_name(),
            total,
            present,
            records.filter(status='absent').count(),
            records.filter(status='late').count(),
            f"{percentage}%",
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename=attendance_{class_obj.name}_{start_date}_{end_date}.xlsx'
    )
    wb.save(response)
    # ws.close() or wb.close()? No need for Response object.
    return response
